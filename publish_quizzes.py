# /// script
# requires-python = ">=3.11"
# dependencies = [
#     "openpyxl",
# ]
# ///
import os
import json
import glob
import subprocess
import openpyxl

def parse_excel_files():
    final_data = {}
    excel_files = [f for f in glob.glob("*.xlsx") if not f.startswith("~$")]

    for file in excel_files:
        print(f"Processing {file}...")
        wb = openpyxl.load_workbook(file, data_only=True)

        intro_data = None
        if 'Introduction' in wb.sheetnames:
            ws_intro = wb['Introduction']
            intro_data = []
            for row in ws_intro.iter_rows(values_only=True):
                clean_row = ["" if cell is None else str(cell) for cell in row]
                # Only add if row is not completely empty
                if any(cell.strip() != "" for cell in clean_row):
                    intro_data.append(clean_row)

        for sheet_name in wb.sheetnames:
            if sheet_name.lower() == 'introduction':
                continue

            ws = wb[sheet_name]
            
            # Get headers from first row
            headers = []
            # Ensure we read the first row properly even if it's empty in places
            for cell in ws[1]:
                headers.append(str(cell.value).strip() if cell is not None and cell.value is not None else "")

            if "Type" not in headers or "Question" not in headers or "Answer" not in headers:
                print(f"  Skipping sheet '{sheet_name}': Missing required headers")
                continue

            type_idx = headers.index("Type")
            q_idx = headers.index("Question")
            ans_idx = headers.index("Answer")

            questions = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                # Check bounds to avoid index errors on empty trailing columns
                if len(row) <= max(type_idx, q_idx, ans_idx):
                    continue
                    
                if not row[type_idx] or not row[q_idx] or not row[ans_idx]:
                    continue # Skip invalid rows

                q_obj = {}
                for i, header in enumerate(headers):
                    if header and i < len(row):
                        val = row[i]
                        # Handle formatting (e.g. floats that should be ints)
                        if isinstance(val, float) and val.is_integer():
                            val = int(val)
                        q_obj[header] = "" if val is None else str(val)

                questions.append(q_obj)

            if questions:
                if sheet_name not in final_data:
                    final_data[sheet_name] = {}
                final_data[sheet_name][sheet_name] = questions
                if intro_data:
                    final_data[sheet_name]["_intro"] = intro_data
                print(f"  Added {len(questions)} questions from sheet '{sheet_name}'.")

    return final_data

def main():
    print("========================================")
    print("   French Quiz App - Mac Admin App      ")
    print("========================================\n")
    
    data = parse_excel_files()

    if not data:
        print("\n❌ No valid quiz data found. Exiting.")
        return

    js_content = "window.quizzesData = " + json.dumps(data, indent=2, ensure_ascii=False) + ";"
    out_path = os.path.join("data", "quizzes.js")

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(js_content)
    print(f"\n✅ Successfully created {out_path}")

    print("\nUploading to GitHub...")
    try:
        # Add the file to git
        subprocess.run(["git", "add", out_path], check=True)
        
        # Check if there are actual changes
        status = subprocess.run(["git", "status", "--porcelain", out_path], capture_output=True, text=True)
        if status.stdout.strip():
            subprocess.run(["git", "commit", "-m", "chore: update quizzes.js via Mac Admin app"], check=True)
            subprocess.run(["git", "push"], check=True)
            print("\n🎉 Successfully pushed to GitHub!")
        else:
            print("\n✅ quizzes.js is already up to date. Nothing to push.")
    except subprocess.CalledProcessError as e:
        print(f"\n❌ Error during Git operations: {e}")

if __name__ == "__main__":
    main()
